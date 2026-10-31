"""Bounded XLSX row parsing: explicit headers, physical rows, no formula execution."""

from datetime import date, datetime, time
from urllib.parse import quote
from xml.etree import ElementTree
from zipfile import ZipFile
import re

from langchain_core.documents import Document
from openpyxl import load_workbook

from app.core.errors import AppError

PARSER_VERSION = "xlsx-row-v1"


def display(value):
    if value is None:
        return "（空）"
    if isinstance(value, bool):
        return "true" if value else "false"
    return value.isoformat() if isinstance(value, (date, datetime, time)) else str(value)


def load_excel(path, settings):
    def reject(message):
        raise AppError("INVALID_EXCEL", message, 422)

    if path.stat().st_size > settings.max_upload_mb * 1024 * 1024:
        reject("Excel 文件超过上传大小限制")

    # Inspect physical XML coordinates too: an inaccurate dimension must not bypass budgets.
    cells = 0
    with ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or any("vbaproject" in name.lower() for name in names):
            reject("不支持重复容器条目或带宏的 Excel")
        for entry in archive.infolist():
            if not entry.filename.startswith("xl/worksheets/") or not entry.filename.endswith(".xml"):
                continue
            with archive.open(entry) as stream:
                for _, node in ElementTree.iterparse(stream, events=("end",)):
                    kind = node.tag.rsplit("}", 1)[-1]
                    if kind == "mergeCell":
                        reject("首版暂不支持合并单元格，请先展开表头和数据")
                    if kind == "row" and not 1 <= int(node.get("r", "0")) <= 10000:
                        reject("Excel 物理行号超过限制")
                    if kind == "c":
                        cells += 1
                        coordinate = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", node.get("r", ""))
                        if not coordinate:
                            reject("Excel 单元格坐标无效")
                        column = 0
                        for character in coordinate[1]:
                            column = column * 26 + ord(character) - 64
                        if int(coordinate[2]) > 10000 or column > 100 or cells > 100000:
                            reject("Excel 超过 10000 行、100 列或总计 100000 单元格限制")
                    node.clear()
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        if len(workbook.worksheets) > 32:
            reject("Excel 工作表超过 32 个")
        documents, characters = [], 0
        budget = max(600, settings.chunk_size)
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            if not 1 <= len(sheet.title) <= 31:
                reject("工作表名称必须为 1 到 31 个字符")
            if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 100:
                reject("Excel 工作表范围超过限制")
            rows = sheet.iter_rows()
            first = next(rows, ())
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in first]
            while headers and not headers[-1]:
                headers.pop()
            if not headers:
                if (sheet.max_row or 0) > 1:
                    reject("每张表的第 1 行必须是表头")
                continue
            if any(not header or len(header) > 100 for header in headers) or len(set(headers)) != len(headers):
                reject("表头不能为空、重复或超过 100 字符")
            if any(cell.data_type == "f" for cell in first):
                reject("表头不能使用公式")
            cached_rows = None
            for row_number, row in enumerate(rows, 2):
                if not any(cell.value is not None for cell in row):
                    continue
                if any(cell.value is not None for cell in row[len(headers):]):
                    reject("数据列缺少对应表头")
                values = []
                for index, cell in enumerate(row[:len(headers)]):
                    value = cell.value
                    if cell.data_type == "f":
                        if cached is None:
                            cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                        # Formula caches are read once in row order, never evaluated or fetched.
                        if cached_rows is None:
                            cached_rows = iter(cached[sheet.title].iter_rows())
                            cached_position = 0
                        while cached_position < row_number:
                            cached_row = next(cached_rows, ())
                            cached_position += 1
                        saved_cell = cached_row[index] if index < len(cached_row) else None
                        if saved_cell is not None and saved_cell.data_type == "e":
                            reject("公式的已保存计算值包含错误，请修正后上传")
                        value = saved_cell.value if saved_cell is not None else None
                        if value is None:
                            reject("公式缺少已保存的计算值，请在 Excel 中计算并保存后上传")
                    if cell.data_type == "e":
                        reject("数据包含 Excel 错误值，请修正后上传")
                    text = display(value)
                    if len(text) > 10000 or "\x00" in text:
                        reject("单元格文本过长或包含无效内容")
                    values.append(text)
                if len(values[0]) > 200:
                    reject("第一列作为行身份，内容不能超过 200 字符")
                prefix = f"工作表：{sheet.title}；{headers[0]}：{values[0]}"
                fragments, current = [], prefix
                for header, value in zip(headers[1:], values[1:]):
                    size = max(1, budget - len(prefix) - len(header) - 4)
                    for offset in range(0, len(value), size):
                        field = f"；{header}：{value[offset:offset + size]}"
                        if len(current) + len(field) > budget:
                            fragments.append(current)
                            current = prefix
                        current += field
                fragments.append(current)
                for content in fragments:
                    characters += len(content)
                    if characters > settings.max_document_characters or len(documents) >= settings.max_document_chunks:
                        reject("Excel 文本或行分片数量超过文档限制")
                    documents.append(Document(page_content=content, metadata={
                        "section": sheet.title, "parser_version": PARSER_VERSION,
                        "excel_source": {"sheet_name": sheet.title, "row_start": row_number, "row_end": row_number,
                                         "source_anchor": f"sheet:{quote(sheet.title, safe='')}:row:{row_number}"},
                    }))
        return documents
    finally:
        workbook.close()
        if cached is not None:
            cached.close()
