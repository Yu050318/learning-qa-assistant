from dataclasses import asdict, replace
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from urllib.parse import quote
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

from app.api.schemas import KnowledgeCitation, V2MessageResponse
from app.application.chat import ChatService
from app.application.evidence import EvidenceRegistry
from app.application.ingestion import IngestionService
from app.application.retriever import RetrieverService
from app.application.sessions import SessionService
from app.core.config import Settings
from app.core.errors import AppError
from app.domain.contracts import AgentChatInput, Chunk, ExcelSource, SearchHit
from app.infrastructure.loaders.local import LocalDocumentLoader
from app.infrastructure.persistence.models import Base, User
from app.infrastructure.vectorstores.milvus import MilvusVectorStore
from app.workflows.supervisor import SupervisorWorkflow
from test_agent_v2 import Models, ScriptedModel, Web, call, final, report


@pytest.fixture
def settings(tmp_path):
    return Settings(_env_file=None, upload_dir=tmp_path / "uploads", parsed_dir=tmp_path / "parsed", embedding_dimension=64)


def workbook(path, rows=None, name="产品参数"):
    book = Workbook()
    sheet = book.active
    sheet.title = name
    for row in rows or [["型号", "额定功率W"], ["NX-020", 180]]:
        sheet.append(row)
    book.save(path)
    book.close()
    return path


def load(path, settings):
    return LocalDocumentLoader(settings).load(path, ".xlsx")


@pytest.mark.parametrize("value,expected", [
    (180, "180"), (12.5, "12.5"), (0, "0"), (False, "false"), (True, "true"),
    (None, "（空）"), (date(2026, 10, 15), "2026-10-15T00:00:00"), ("中文单位 kg", "中文单位 kg"),
])
def test_values_keep_header_and_units(tmp_path, settings, value, expected):
    path = workbook(tmp_path / "input.xlsx", [["型号", "功率W"], ["NX-020", value]])
    doc = load(path, settings)[0]
    assert f"功率W：{expected}" in doc.page_content
    assert doc.metadata["excel_source"]["row_start"] == 2


def test_multisheet_and_blank_rows_preserve_physical_location(tmp_path, settings):
    path = workbook(tmp_path / "input.xlsx", [["型号", "功率W"], ["NX-020", 180], [], ["NX-020", 180]])
    from openpyxl import load_workbook
    book = load_workbook(path)
    other = book.create_sheet("库存 / 不允许".replace("/", "／"))
    other.append(["型号", "库存"])
    other.append(["NX-020", 42])
    hidden = book.create_sheet("隐藏")
    hidden.append(["型号", "秘密"])
    hidden.append(["NX-020", "不索引"])
    hidden.sheet_state = "hidden"
    book.save(path)
    book.close()
    docs = load(path, settings)
    assert [doc.metadata["excel_source"]["row_start"] for doc in docs] == [2, 4, 2]
    assert len({doc.metadata["excel_source"]["source_anchor"] for doc in docs}) == 3
    assert "不索引" not in str(docs)


@pytest.mark.parametrize("rows", [
    [["型号", "型号"], ["a", 1]],
    [[None, "功率"], ["a", 1]],
    [["型号"], ["a", 1]],
    [["x" * 101], ["a"]],
    [["=1+1", "功率"], ["a", 1]],
    [[None, None], ["a", 1]],
])
def test_bad_headers_rejected(tmp_path, settings, rows):
    with pytest.raises(AppError) as error:
        load(workbook(tmp_path / "input.xlsx", rows), settings)
    assert error.value.code == "INVALID_EXCEL"


def test_uncached_formula_is_rejected_without_execution(tmp_path, settings):
    with pytest.raises(AppError, match="公式缺少"):
        load(workbook(tmp_path / "input.xlsx", [["型号", "计算"], ["a", "=1+1"]]), settings)


def test_cached_formula_uses_saved_value(tmp_path, settings):
    path = workbook(tmp_path / "input.xlsx", [["型号", "计算"], ["a", "=1+1"]])
    with ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    assert "计算：2" in load(path, settings)[0].page_content


@pytest.mark.parametrize("change", ["merge", "row_budget", "column_budget", "long_cell", "long_identity", "sheet_name", "error"])
def test_excel_boundaries(tmp_path, settings, change):
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["型号", "值"])
    sheet.append(["a", 42])
    if change == "merge":
        sheet.merge_cells("A1:B1")
    elif change == "row_budget":
        sheet.cell(10001, 1, "a")
    elif change == "column_budget":
        sheet.cell(2, 101, "a")
    elif change == "long_cell":
        sheet.cell(2, 2, "a" * 10001)
    elif change == "long_identity":
        sheet.cell(2, 1, "a" * 201)
    elif change == "sheet_name":
        sheet.title = "x" * 32
    else:
        sheet.cell(2, 2, "#DIV/0!")
    book.save(path)
    book.close()
    with pytest.raises(AppError):
        load(path, settings)


def test_long_rows_keep_identity_in_every_fragment_and_stable_ids(tmp_path, settings):
    path = workbook(tmp_path / "input.xlsx", [["型号", "说明"], ["NX-020", "文字" * 1000]])
    loader = LocalDocumentLoader(settings)
    docs = load(path, settings)
    first = loader.chunks(docs, uuid4(), uuid4(), "input.xlsx")
    assert len(first) > 1
    assert all("型号：NX-020" in chunk.content and len(chunk.content) <= settings.chunk_size for chunk in first)
    assert len({chunk.excel_source.source_anchor for chunk in first}) == 1
    assert all(chunk.page_number is None for chunk in first)
    user_id, doc_id = uuid4(), uuid4()
    assert loader.chunks(docs, user_id, doc_id, "input.xlsx") == loader.chunks(load(path, settings), user_id, doc_id, "input.xlsx")


def test_document_budget(tmp_path, settings):
    settings.max_document_chunks = 1
    path = workbook(tmp_path / "input.xlsx", [["型号", "值"], ["a", 1], ["b", 2]])
    with pytest.raises(AppError):
        load(path, settings)


def source(row=2):
    return ExcelSource("参数", row, row, f"sheet:{quote('参数', safe='')}:row:{row}")


def test_retrieval_does_not_collapse_identical_values_on_different_rows(settings):
    uid, did = uuid4(), uuid4()
    base = Chunk(str(uuid4()), str(uid), str(did), 0, "相同文本", "input.xlsx", excel_source=source())
    hits = [SearchHit(base, .03), SearchHit(replace(base, chunk_id=str(uuid4()), excel_source=source(4)), .02)]
    vectors = SimpleNamespace(search=lambda *args: hits)
    embedding = SimpleNamespace(embed_query=lambda query: [0.] * 64)
    result = RetrieverService(embedding, vectors, 6).retrieve(uid, "问题", [did])
    assert len(result) == 2
    registry = EvidenceRegistry()
    tool_result = registry.add_knowledge(result)
    assert tool_result["sources"][0]["excel_source"]["row_start"] == 2
    citations = registry.validate("grounded", "两个位置[1][2]")
    assert [KnowledgeCitation.model_validate(value).excel_source.row_start for value in citations] == [2, 4]


def test_milvus_roundtrip_and_legacy_collection(settings):
    store = MilvusVectorStore(settings)
    uid, did = uuid4(), uuid4()
    chunk = Chunk(str(uuid4()), str(uid), str(did), 0, "额定功率：180", "input.xlsx", excel_source=source())
    client = MagicMock()
    connection = MagicMock()
    connection.__enter__.return_value = client
    with patch.object(store, "connection", return_value=connection), patch.object(store, "_validate", return_value={"excel_source"}):
        store.upsert([chunk], [[0.] * 64])
        row = client.upsert.call_args.kwargs["data"][0]
        assert row["excel_source"] == asdict(source())
        entity = {key: value for key, value in row.items() if key not in {"vector", "embedding_model"}}
        client.hybrid_search.return_value = [[{"entity": entity, "distance": .03}]]
        assert store.search(uid, [did], "功率", [0.] * 64, 6)[0].chunk == chunk
    with patch.object(store, "connection", return_value=connection), patch.object(store, "_validate", return_value=set()):
        with pytest.raises(AppError) as error:
            store.upsert([chunk], [[0.] * 64])
        assert error.value.code == "VECTOR_PROVENANCE_REQUIRED"
        store.upsert([replace(chunk, excel_source=None)], [[0.] * 64])
        assert "excel_source" not in client.upsert.call_args.kwargs["data"][0]


def test_additive_upgrade_is_explicit_and_idempotent(settings):
    store = MilvusVectorStore(settings)
    client = MagicMock()
    connection = MagicMock()
    connection.__enter__.return_value = client
    with patch.object(store, "connection", return_value=connection), patch.object(store, "_validate", side_effect=[set(), {"excel_source"}, {"excel_source"}]):
        store.add_excel_provenance()
        store.add_excel_provenance()
    client.add_collection_field.assert_called_once()
    client.drop_collection.assert_not_called()
    client.delete.assert_not_called()


def test_bad_vector_provenance_rejected_and_old_search_has_no_new_field(settings):
    store = MilvusVectorStore(settings)
    uid, did = uuid4(), uuid4()
    chunk = Chunk(str(uuid4()), str(uid), str(did), 0, "旧文档内容", "old.md")
    client = MagicMock()
    connection = MagicMock()
    connection.__enter__.return_value = client
    entity = {**asdict(chunk), "page_number": 0, "section": ""}
    entity.pop("excel_source")
    with patch.object(store, "connection", return_value=connection), patch.object(store, "_validate", return_value=set()):
        client.hybrid_search.return_value = [[{"entity": dict(entity), "distance": .03}]]
        result = store.search(uid, [did], "内容", [0.] * 64, 6)
        assert result[0].chunk.excel_source is None
        assert "excel_source" not in client.hybrid_search.call_args.kwargs["output_fields"]
    with patch.object(store, "connection", return_value=connection), patch.object(store, "_validate", return_value={"excel_source"}):
        client.hybrid_search.return_value = [[{"entity": {**entity, "excel_source": {**asdict(source()), "source_anchor": "伪造来源"}}, "distance": .03}]]
        with pytest.raises(AppError) as error:
            store.search(uid, [did], "内容", [0.] * 64, 6)
        assert error.value.code == "VECTOR_METADATA_INVALID"


def test_upload_to_supervisor_to_persisted_citation(tmp_path, settings):
    engine = create_engine("sqlite://")
    connection = engine.connect()
    connection.execute(text("ATTACH DATABASE ':memory:' AS rag"))
    Base.metadata.create_all(connection)
    database = SimpleNamespace(sessions=sessionmaker(bind=connection, expire_on_commit=False))
    with database.sessions() as session:
        user = User(username="excel-reader", password_hash="fixture")
        session.add(user)
        session.commit()
        user_id = user.id

    class Vectors:
        def __init__(self):
            self.chunks = []

        def delete_document(self, uid, did):
            self.chunks = [chunk for chunk in self.chunks if chunk.user_id != str(uid) or chunk.document_id != str(did)]

        def upsert(self, chunks, vectors):
            self.chunks.extend(chunks)

        def search(self, uid, dids, *args):
            return [SearchHit(chunk, .03) for chunk in self.chunks if chunk.user_id == str(uid) and chunk.document_id in {str(did) for did in dids}]

    vectors = Vectors()
    embeddings = SimpleNamespace(embed_documents=lambda values: [[0.] * 64 for _ in values], embed_query=lambda value: [0.] * 64)
    loader = LocalDocumentLoader(settings)
    ingestion = IngestionService(settings, database, loader, None, embeddings, vectors)
    try:
        path = workbook(tmp_path / "input.xlsx")
        document = ingestion.upload(user_id, "input.xlsx", None, BytesIO(path.read_bytes()))
        assert document.parser_metadata["provider"] == "local"
        ingestion.process(user_id, document.id)
        assert ingestion.get(user_id, document.id).status == "ready"
        models = Models(ScriptedModel(responses=[call("delegate_rag", task="查询功率"), call("search_knowledge", query="NX-020"), report(), final("额定功率 180W[1]")]))
        workflow = SupervisorWorkflow(settings, RetrieverService(embeddings, vectors, 6), models, Web())
        sessions = SessionService(database)
        chat = ChatService(settings, database, sessions, workflow)
        conversation = sessions.create(user_id, "表格问答")
        events = []
        message = chat.answer_v2(user_id, conversation.id, AgentChatInput("NX-020 的功率？", search_mode="knowledge"), "excel-test", lambda event, value: events.append(event))
        assert message.citations[0]["excel_source"]["row_start"] == 2
        _, history = sessions.get(user_id, conversation.id)
        persisted = V2MessageResponse.model_validate(history[-1])
        assert persisted.citations[0].excel_source.sheet_name == "产品参数"
        assert "answer_delta" in events
        with pytest.raises(AppError):
            ingestion.get(uuid4(), document.id)
    finally:
        ingestion.close()
        connection.close()
        engine.dispose()
