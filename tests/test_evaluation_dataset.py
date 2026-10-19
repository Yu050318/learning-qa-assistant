import json
from collections import Counter
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.application.evaluation_dataset import GoldAnchor, WorkbenchDataset
from scripts.verify_enterprise_fixture import verify

FIXTURE = Path(__file__).resolve().parents[1] / "evals/rag/synthetic_enterprise"


def test_enterprise_assets_hashes_labels_and_real_spreadsheet_anchors():
    result = verify(FIXTURE)
    assert (result["documents"], result["cases"], result["dev"], result["frozen"]) == (12, 80, 20, 60)
    dataset = WorkbenchDataset.model_validate_json((FIXTURE / "dataset.json").read_text(encoding="utf-8"))
    assert Counter(case.category for case in dataset.cases) == {"simple_fact": 20, "table_query": 20, "cross_document": 20, "no_answer": 10, "version_policy": 10}
    for name, count in (("product_specs.xlsx", 200), ("product_prices.xlsx", 200), ("inventory.xlsx", 200), ("repair_records.xlsx", 100)):
        workbook = load_workbook(FIXTURE / "upload_corpus" / name, read_only=True, data_only=True)
        try:
            assert workbook.active.max_row == count + 1
            for case in dataset.cases:
                for anchor in case.gold_anchors:
                    if anchor.canonical_name == name:
                        sheet = workbook[anchor.sheet]
                        assert 2 <= anchor.row_start <= anchor.row_end <= sheet.max_row
                        values = list(sheet.iter_rows(min_row=anchor.row_start, max_row=anchor.row_end, values_only=True))
                        assert any(any(str(value).startswith("NX-") for value in row) for row in values)
        finally:
            workbook.close()


@pytest.mark.parametrize("field,value", [("user_id", str(uuid4())), ("document_ids", [str(uuid4())]), ("web_urls", ["https://example.com"]), ("search_mode", "web")])
def test_labels_cannot_choose_execution_identity_private_scope_or_external_access(field, value):
    data = json.loads((FIXTURE / "dataset.json").read_text(encoding="utf-8"))
    data["cases"][0][field] = value
    with pytest.raises(ValidationError):
        WorkbenchDataset.model_validate(data)


def test_duplicate_cases_and_partial_row_anchors_are_rejected():
    data = json.loads((FIXTURE / "dataset.json").read_text(encoding="utf-8"))
    data["cases"].append(data["cases"][0])
    with pytest.raises(ValidationError, match="unique"):
        WorkbenchDataset.model_validate(data)
    for fields in ({"row_start": 2}, {"row_start": 3, "row_end": 2, "sheet": "Sheet1"}, {"row_start": 2, "row_end": 3}):
        with pytest.raises(ValidationError):
            GoldAnchor(canonical_name="inventory.xlsx", **fields)
    with pytest.raises(ValidationError):
        GoldAnchor(canonical_name="../dataset.json")


def test_frozen_hash_check_rejects_changed_labels(tmp_path):
    import shutil
    shutil.copytree(FIXTURE, tmp_path / "fixture")
    target = tmp_path / "fixture" / "dataset.json"
    target.write_text(target.read_text(encoding="utf-8").replace("NX-041", "NX-999"), encoding="utf-8")
    with pytest.raises(ValueError, match="hash mismatch"):
        verify(target.parent)
